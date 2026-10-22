from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'MedSafe_AI_Complete_Project_Guide.docx'
ASSET = ROOT / 'docs' / '_guide_assets'
ASSET.mkdir(exist_ok=True)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def borders(table, color='D9E2E8'):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = b.find(qn('w:'+edge))
        if e is None: e = OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),color)

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'0B3B4A')
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val); shade(cells[i], 'F4F8FA' if ridx%2 else 'FFFFFF')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def diagram(path):
    W,H=1500,650; im=Image.new('RGB',(W,H),'#F4F8FA'); d=ImageDraw.Draw(im)
    try: font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28); small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
    except: font=small=None
    boxes=[('1  Upload','Prescription image or PDF'),('2  Read','Tesseract extracts text'),('3  Check','Person corrects the text'),('4  Review','Fields and safety findings'),('5  Report','Saved local report')]
    x0=55; bw=250; gap=45
    for i,(title,sub) in enumerate(boxes):
        x=x0+i*(bw+gap); d.rounded_rectangle((x,190,x+bw,420),radius=22,fill='#FFFFFF',outline='#11B6A2',width=5)
        d.text((x+22,225),title,fill='#0B3B4A',font=font); d.multiline_text((x+22,285),sub,fill='#4B6873',font=small,spacing=8)
        if i<len(boxes)-1:
            d.line((x+bw+8,305,x+bw+gap-8,305),fill='#11B6A2',width=7); d.polygon([(x+bw+gap-8,305),(x+bw+gap-28,292),(x+bw+gap-28,318)],fill='#11B6A2')
    im.save(path)

def model_diagram(path):
    W,H=1400,700; im=Image.new('RGB',(W,H),'#FFFFFF'); d=ImageDraw.Draw(im)
    try: f=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28); s=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',21)
    except: f=s=None
    items=[('Input image','PNG / JPG / PDF'),('CRAFT detector','Finds text regions'),('Local research model','TrOCR checkpoint'),('Decoded text','Unverified transcript'),('Human review','Required before report')]
    for i,(a,b) in enumerate(items):
        y=55+i*125; d.rounded_rectangle((170,y,1230,y+85),radius=18,fill='#F4F8FA',outline='#0B3B4A',width=3); d.text((210,y+15),a,fill='#0B3B4A',font=f); d.text((700,y+22),b,fill='#4B6873',font=s)
        if i<len(items)-1: d.line((700,y+85,700,y+120),fill='#11B6A2',width=6); d.polygon([(700,y+120),(688,y+100),(712,y+100)],fill='#11B6A2')
    im.save(path)

def add_img(doc, path, caption, width=6.2):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic=True; c.runs[0].font.size=Pt(9); c.runs[0].font.color.rgb=RGBColor(80,95,100)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size,color in [('Title',29,'0B3B4A'),('Heading 1',19,'0B3B4A'),('Heading 2',14,'0E756D'),('Heading 3',11,'0B3B4A')]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)

title=doc.add_paragraph(style='Title'); title.add_run('MedSafe AI Complete Project Guide'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Application, research model, Drive archive, workflow, testing, and thesis evidence'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor(14,117,109)
p=doc.add_paragraph('Prepared for project review | 22 October 2026'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('This guide explains what the project is, how the folders relate to one another, how a prescription moves through the application, where the model files come from, what has been tested, and what must still be stated carefully in a thesis. It is written for a reader who may not know the internal software terminology.')

doc.add_heading('1 Project in one minute',1)
doc.add_paragraph('MedSafe AI is a local research application that reads a prescription image, presents the extracted text for human correction, organizes the available prescription details, and produces a review report. It is designed to support research and careful review. It is not a prescribing system and it does not automatically approve a medicine or dosage.')
table(doc,['Question','Answer'],[['What does it read?','Prescription images and supported single-page documents.'],['Where is information stored?','A local SQLite database and local files under runtime_data.'],['What is the normal reader?','Tesseract, selected for the ordinary product workflow.'],['What is the research model?','A locally stored fine-tuned TrOCR checkpoint used only when explicitly configured.'],['What is always required?','Human checking of the extracted text before a report is treated as reviewed.']])

doc.add_heading('2 The complete workflow',1)
flow=ASSET/'workflow.png'; diagram(flow); add_img(doc,flow,'Figure 1. The end-to-end prescription workflow.')
doc.add_paragraph('The important design decision is that reading and reviewing are separate steps. The application does not treat an OCR result as a confirmed medical instruction. The person reviewing the record must compare the text with the original image, correct errors, and confirm the transcription before the report is generated.')

doc.add_heading('3 What each part of the project does',1)
table(doc,['Folder or file','Purpose'],[
['medsafe/','Main application logic: reading images, extracting fields, storing records, validating information, and producing reports.'],
['backend/','A compatibility server layer and organized API services. It keeps the API entry point separate from the core logic.'],
['frontend/','The React web interface used for upload, review, history, and reports.'],
['tests/','Automated checks for the API, OCR behavior, review rules, reports, history, and model integration.'],
['02_models/','Local research model files. These are intentionally not uploaded to GitHub because they are large.'],
['model_cache/','Local detector files used by optional image-reading components.'],
['runtime_data/','Local database, uploaded images, generated reports, and verification output.'],
['examples/','Safe demonstration image and matching example text.'],
['scripts/','Setup and verification helpers.'],
['docs/','Human-readable project, evaluation, and research documentation.'],
['start_product.py and start_windows.bat','Startup helpers for local use.'],
['requirements*.txt','Lists of packages for the ordinary application and optional research components.'],
])

doc.add_heading('4 The Drive project folder',1)
doc.add_paragraph('The Google Drive project is the research archive. It contains the material created during the research phase: notebooks, model checkpoints, experiment outputs, datasets, and supporting files. The local application is the runnable product copy; the Drive folder is the source archive used to recover research material and verify where a model came from.')
table(doc,['Drive area','How it relates to the application'],[
['02_models','Contains the fine-tuned model checkpoint. A selected checkpoint is copied or mounted locally before model verification.'],
['Research notebooks','Explain the experiments, image preparation, detector settings, recognition settings, and saved outputs.'],
['Datasets and labels','Needed to calculate genuine accuracy and compare predictions with known correct transcriptions.'],
['Experiment results','Used to support the thesis discussion, but not automatically treated as application proof.'],
['Local project folder','Contains the runnable code, tests, local database, and documentation used for the final demonstration.'],
])
doc.add_paragraph('The Drive archive and the GitHub repository serve different purposes. GitHub contains source code and safe examples. The Drive archive holds large or sensitive research artifacts. The local machine joins them when the application needs the model or research data.')

doc.add_heading('5 How the model works',1)
md=ASSET/'model_pipeline.png'; model_diagram(md); add_img(doc,md,'Figure 2. Model and human-review stages.')
doc.add_paragraph('The model path is deliberately local. The application checks that the expected configuration, processor, tokenizer, and weight files exist before loading the model. It uses local-only loading and does not silently download a different public model when the selected research checkpoint is missing.')
table(doc,['Model detail','Current project information'],[
['Image detector','CRAFT through the local EasyOCR model cache for the optional research path.'],
['Text recognizer','Fine-tuned TrOCR checkpoint selected from the research project.'],
['Loading behavior','Local files only; missing files produce an explicit setup error.'],
['Device behavior','Uses the available graphics processor when available; otherwise uses the processor.'],
['Output meaning','A proposed transcript, not a medically confirmed instruction.'],
['Verification evidence','The stored smoke-test report records detector loading, model loading, inference, region count, and warnings.'],
])
doc.add_paragraph('The current verification report records a successful model load and inference run with five detected regions and five non-empty predictions. It also correctly records that text and confidence values remain unverified. This is evidence that the pipeline runs, not evidence that the handwriting was read perfectly.')

doc.add_heading('6 Screens from the application',1)
imgs=sorted([p for p in (ROOT/'runtime_data').glob('*') if p.suffix.lower() in ('.png','.jpg','.jpeg')], key=lambda p:p.name)
for pth in imgs[:6]:
    add_img(doc,pth,f'Figure. Application or demonstration image retained in the local project: {pth.name}.',5.6)

doc.add_heading('7 Data and report flow',1)
doc.add_paragraph('An upload creates a local record. The original image is retained, the extracted transcript is stored, and later corrections are recorded as a new review state. This makes it possible to compare the original reading with the corrected text. The generated report contains findings and uncertainty indicators; it does not grant permission to dispense medication.')
table(doc,['Stage','Stored or shown'],[['Upload','Original file, filename, type, and local record identifier.'],['Reading','Extracted text, regions where available, engine name, and warnings.'],['Review','Corrected transcription and confirmation action.'],['Assessment','Extracted fields, missing information, and rule-based findings.'],['History','Previous records and their current review state.'],['Export','A JSON report for research review and record keeping.']])

doc.add_heading('8 Testing and current evidence',1)
table(doc,['Check','Result'],[['Automated backend suite','109 tests passed on 22 October 2026.'],['Frontend production build','Successful.'],['Local research-model smoke test','Passed model initialization and image inference.'],['Local database','Present under runtime_data.'],['GitHub repository','Created and pushed to the main branch.'],['Clinical validation','Not provided and must not be claimed.']])
doc.add_paragraph('The automated tests demonstrate that the software behaves as designed for its fixtures and workflow checks. They do not measure real-world handwriting accuracy. That separate measurement requires human-labelled prescription images and correct reference transcriptions.')

doc.add_heading('9 What the thesis can claim',1)
doc.add_paragraph('The thesis can accurately claim that a local prescription-reading and review workflow was implemented, that Tesseract is integrated as the ordinary reader, that the research model can be loaded locally when its checkpoint is available, and that the application has automated workflow tests. It should not claim perfect recognition, clinical safety, automatic dispensing approval, or medical diagnosis.')
table(doc,['Evidence level','Safe wording'],[['Software behavior','The application supports upload, extraction, correction, review, history, and report generation.'],['Model readiness','The selected local checkpoint loaded and produced non-empty output on the smoke-test image.'],['Recognition quality','Accuracy must be reported only after evaluation against labelled test images.'],['Safety findings','The system reports unresolved or rule-based findings; it does not establish clinical safety.']])

doc.add_heading('10 Remaining research evidence',1)
doc.add_paragraph('The main remaining thesis gap is not application startup. It is measured evidence. A suitable evaluation set should contain unique images, human-verified transcripts, structured fields, and independently reviewed alert labels. The benchmark should report character error rate, word error rate, field accuracy, and processing time for the same images and preprocessing settings.')
table(doc,['Needed item','Why it matters'],[['Human-verified images','Provides a trustworthy test sample.'],['Correct transcripts','Allows recognition errors to be counted.'],['Structured field labels','Allows medicine, dose, frequency, age, and route to be measured separately.'],['Independent rule review','Prevents invented or unsupported safety conclusions.'],['Screenshots and run logs','Shows that the final application actually performs the described workflow.']])

doc.add_heading('11 How to start the application',1)
doc.add_paragraph('On Windows, open PowerShell in the project folder and run start_windows.bat. The script prepares the local environment, builds the web interface, and starts the server at http://127.0.0.1:8000. For a direct verification, run python -m pytest -q. The research model is optional for the ordinary Tesseract workflow and must be configured from its local checkpoint directory before running the model verification command.')
doc.add_heading('12 Important limitations',1)
for text in [
    'This is a local research application, not a licensed clinical device.',
    'OCR can misread medicine names, numbers, units, and instructions, especially handwriting.',
    'The local database and uploaded files are not encrypted at rest by this application.',
    'The application should not be exposed to the public internet without authentication, access control, and a formal security review.',
    'The model checkpoint is large and intentionally kept outside the GitHub source repository.',
    'Human review remains mandatory before using any extracted text for a research report.'
]: doc.add_paragraph(text, style='List Bullet')
doc.add_heading('13 Final project assessment',1)
doc.add_paragraph('The software product is operational and test-covered. The research contribution becomes thesis-ready when the selected checkpoint is confirmed against the original notebook, the labelled evaluation set is run, the results are interpreted honestly, and the screenshots and run logs are attached as evidence. The project already provides the structure needed to complete that final research stage.')

doc.sections[0].footer.paragraphs[0].text='MedSafe AI | Complete Project Guide | Local research use only'
doc.save(OUT)
print(OUT)
